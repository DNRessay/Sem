"""Uploaded files → something the model can use. Documents become text (PDF, Word, Excel, PowerPoint,
OpenDocument, RTF, EPUB, old .doc), audio becomes a transcript (Groq Whisper, Gemini as fallback), video
and scanned PDFs are read by Gemini, and any image format becomes a JPEG the vision model takes."""
import base64
import html
import io
import re
import zipfile

import httpx

from config import settings
from tools import gemini_media

MAX_TEXT = 60_000
VISION_MIMES = {"image/png", "image/jpeg", "image/webp", "image/gif"}
AUDIO_EXT = {"mp3", "wav", "m4a", "aac", "ogg", "oga", "opus", "flac", "webm", "weba", "amr", "aiff", "aif", "mp4a", "3gp"}
VIDEO_EXT = {"mp4", "mov", "m4v", "avi", "mkv", "3gpp", "mpeg", "mpg", "wmv"}
_GEMINI_AUDIO = {"mp3": "audio/mp3", "wav": "audio/wav", "aac": "audio/aac", "m4a": "audio/aac", "ogg": "audio/ogg",
                 "oga": "audio/ogg", "opus": "audio/ogg", "flac": "audio/flac", "aiff": "audio/aiff", "aif": "audio/aiff"}
_GEMINI_VIDEO = {"mp4": "video/mp4", "m4v": "video/mp4", "mov": "video/quicktime", "avi": "video/x-msvideo",
                 "mpeg": "video/mpeg", "mpg": "video/mpeg", "wmv": "video/x-ms-wmv", "3gpp": "video/3gpp", "mkv": "video/mp4"}


def ext_of(name: str) -> str:
    return (name.rsplit(".", 1)[-1] if "." in name else "").lower()


def is_image(name: str, mime: str) -> bool:
    return mime.startswith("image/") and ext_of(name) != "svg" and mime != "image/svg+xml"


def to_vision_image(raw: bytes, mime: str) -> tuple[bytes, str] | None:
    """HEIC, BMP, TIFF, AVIF, ICO… → JPEG (the vision model takes PNG/JPEG/WEBP/GIF only). None if unreadable."""
    if mime in VISION_MIMES:
        return raw, mime
    try:
        from PIL import Image
        try:
            import pillow_heif
            pillow_heif.register_heif_opener()
        except ImportError:
            pass
        img = Image.open(io.BytesIO(raw))
        img.seek(0)
        img = img.convert("RGB")
        img.thumbnail((2048, 2048))
        out = io.BytesIO()
        img.save(out, "JPEG", quality=88)
        return out.getvalue(), "image/jpeg"
    except Exception:
        return None


def _clip(text: str) -> str:
    text = re.sub(r"\n{3,}", "\n\n", (text or "").replace("\r", "")).strip()
    return text if len(text) <= MAX_TEXT else text[:MAX_TEXT] + "\n…(cut — the file is longer)"


def _xml_text(xml: str) -> str:
    xml = re.sub(r"(?=<(?:text:p|text:h|w:p|a:p|p|br|div|li|h\d|tr)[\s>/])", "\n", xml)
    xml = re.sub(r"(?=<(?:table:table-cell|w:tab|td|th)[\s>/])", "\t", xml)
    return html.unescape(re.sub(r"<[^>]+>", "", xml))


def _pdf(raw: bytes) -> str:
    from pypdf import PdfReader
    return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(raw)).pages)


def _docx(raw: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(raw))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        parts += ["\t".join(c.text for c in row.cells) for row in t.rows]
    return "\n".join(parts)


def _xlsx(raw: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                out.append(",".join("" if v is None else str(v) for v in row))
    return "\n".join(out)


def _xls(raw: bytes) -> str:
    import xlrd
    book = xlrd.open_workbook(file_contents=raw)
    out = []
    for sh in book.sheets():
        out.append(f"## Sheet: {sh.name}")
        out += [",".join(str(c.value) for c in sh.row(r)) for r in range(sh.nrows)]
    return "\n".join(out)


def _pptx(raw: bytes) -> str:
    from pptx import Presentation
    out = []
    for i, slide in enumerate(Presentation(io.BytesIO(raw)).slides, 1):
        out.append(f"## Slide {i}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                out += ["\t".join(c.text for c in row.cells) for row in shape.table.rows]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append(f"Notes: {slide.notes_slide.notes_text_frame.text}")
    return "\n".join(out)


def _zip_xml(raw: bytes, names) -> str:
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        wanted = [n for n in z.namelist() if names(n)]
        return "\n".join(_xml_text(z.read(n).decode("utf-8", "ignore")) for n in sorted(wanted))


def _rtf(raw: bytes) -> str:
    from striprtf.striprtf import rtf_to_text
    return rtf_to_text(raw.decode("latin-1"))


def _doc(raw: bytes) -> str:
    """Old binary Word: its text is stored as UTF-16 (or 8-bit) runs — pull out the readable ones."""
    runs = re.findall(r"(?:[\x20-\x7e -ɏ\r\n\t][\x00]){4,}", raw.decode("latin-1"))
    text = "\n".join(r.replace("\x00", "") for r in runs)
    if len(text) < 40:
        plain = "\n".join(re.findall(r"[\x20-\x7e\r\n\t]{6,}", raw.decode("latin-1")))
        text = plain if len(plain) > len(text) else text
    return text


_DOCS = {
    "pdf": _pdf, "docx": _docx, "docm": _docx, "dotx": _docx, "xlsx": _xlsx, "xlsm": _xlsx, "xls": _xls,
    "pptx": _pptx, "ppsx": _pptx, "rtf": _rtf, "doc": _doc,
    "odt": lambda r: _zip_xml(r, lambda n: n == "content.xml"),
    "ods": lambda r: _zip_xml(r, lambda n: n == "content.xml"),
    "odp": lambda r: _zip_xml(r, lambda n: n == "content.xml"),
    "epub": lambda r: _zip_xml(r, lambda n: n.endswith((".xhtml", ".html", ".htm"))),
    "pages": lambda r: _zip_xml(r, lambda n: n.endswith(".xml")),
}
_DOC_MIMES = {
    "application/pdf": "pdf", "application/msword": "doc", "application/rtf": "rtf", "text/rtf": "rtf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.ms-excel": "xls",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "application/vnd.oasis.opendocument.text": "odt", "application/vnd.oasis.opendocument.spreadsheet": "ods",
    "application/vnd.oasis.opendocument.presentation": "odp", "application/epub+zip": "epub",
}


async def _groq_transcribe(raw: bytes, name: str) -> str:
    if not settings.GROQ_API_KEY:
        return ""
    try:
        async with httpx.AsyncClient(timeout=120) as client:
            r = await client.post("https://api.groq.com/openai/v1/audio/transcriptions",
                                  headers={"Authorization": f"Bearer {settings.GROQ_API_KEY}"},
                                  files={"file": (name, raw)}, data={"model": "whisper-large-v3-turbo"})
        return r.json().get("text", "") if r.status_code == 200 else ""
    except (httpx.HTTPError, ValueError):
        return ""


async def _gemini_read(raw: bytes, mime: str, ask: str) -> str:
    result = await gemini_media._generate(settings.GEMINI_MODEL, {"contents": [{"parts": [
        {"inlineData": {"mimeType": mime, "data": base64.b64encode(raw).decode()}}, {"text": ask}]}]}, timeout=120)
    if not result["ok"]:
        return ""
    parts = (((result["data"].get("candidates") or [{}])[0].get("content") or {}).get("parts")) or []
    return "".join(p.get("text", "") for p in parts)


async def read_file(name: str, mime: str, raw: bytes) -> tuple[str, str]:
    """(text, how) for a non-image upload; text is "" when nothing could be read, with `how` saying why."""
    ext = ext_of(name) or _DOC_MIMES.get(mime, "")
    kind = _DOC_MIMES.get(mime) if ext not in _DOCS else ext
    if kind in _DOCS:
        try:
            text = _DOCS[kind](raw)
        except Exception as e:
            text, err = "", str(e)[:150]
        else:
            err = ""
        if kind == "pdf" and len(text.strip()) < 50:  # a scan: let Gemini read the pages
            text = await _gemini_read(raw, "application/pdf", "Transcribe all the text in this PDF, keeping its structure. "
                                                              "Describe any charts, tables or images briefly.")
            return _clip(text), "read by Gemini (scanned PDF)" if text else "no text found in this PDF"
        return _clip(text), (f"couldn't read it: {err}" if err else "") if not text.strip() else "text extracted"
    if ext in AUDIO_EXT or mime.startswith("audio/"):
        text = await _groq_transcribe(raw, name or f"audio.{ext or 'mp3'}")
        how = "transcribed (Groq Whisper)"
        if not text:
            text = await _gemini_read(raw, _GEMINI_AUDIO.get(ext, mime or "audio/mp3"), "Transcribe this audio word for word.")
            how = "transcribed (Gemini)"
        return _clip(text), how if text else "couldn't transcribe the audio"
    if ext in VIDEO_EXT or mime.startswith("video/"):
        text = await _gemini_read(raw, _GEMINI_VIDEO.get(ext, mime or "video/mp4"),
                                  "Describe this video scene by scene and transcribe any speech and on-screen text.")
        return _clip(text), "watched by Gemini" if text else "couldn't read the video"
    if ext in ("gdoc", "gsheet", "gslides"):
        return "", "a Google Drive shortcut, not the file — download it as Word/Excel/PowerPoint, or ask Sem to open it from Drive"
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        return "", f"unsupported file type ({mime or ext or 'unknown'})"
    return _clip(text), "read as text"
