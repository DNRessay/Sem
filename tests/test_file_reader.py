import io
import zipfile

import pytest

from tools import file_reader


def _docx():
    import docx
    d = docx.Document()
    d.add_paragraph("Quarterly report")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Revenue", "R120k"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx():
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "Sales"
    wb.active.append(["Month", "Total"])
    wb.active.append(["Jan", 4500])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pptx():
    from pptx import Presentation
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Launch plan"
    buf = io.BytesIO()
    p.save(buf)
    return buf.getvalue()


def _odt():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("content.xml", '<office:document-content><text:p>Hello from LibreOffice &amp; co</text:p></office:document-content>')
    return buf.getvalue()


@pytest.mark.asyncio
@pytest.mark.parametrize("name,raw,expect", [
    ("report.docx", _docx(), ["Quarterly report", "Revenue\tR120k"]),
    ("sales.xlsx", _xlsx(), ["## Sheet: Sales", "Jan,4500"]),
    ("deck.pptx", _pptx(), ["## Slide 1", "Launch plan"]),
    ("notes.odt", _odt(), ["Hello from LibreOffice & co"]),
    ("memo.rtf", rb"{\rtf1\ansi Hello {\b bold} world}", ["Hello bold world"]),
    ("old.doc", "Minutes of the meeting held on Monday".encode("utf-16-le"), ["Minutes of the meeting"]),
])
async def test_documents_become_text(name, raw, expect):
    text, how = await file_reader.read_file(name, "", raw)
    assert all(e in text for e in expect), text
    assert how == "text extracted"


@pytest.mark.asyncio
async def test_audio_is_transcribed_by_groq_then_gemini(monkeypatch):
    async def groq(raw, name):
        return ""

    async def gemini(raw, mime, ask):
        assert mime == "audio/aac"
        return "hello sem"

    monkeypatch.setattr(file_reader, "_groq_transcribe", groq)
    monkeypatch.setattr(file_reader, "_gemini_read", gemini)
    assert await file_reader.read_file("voice.m4a", "audio/mp4", b"x") == ("hello sem", "transcribed (Gemini)")


@pytest.mark.asyncio
async def test_unreadable_files_say_why():
    text, how = await file_reader.read_file("blob.bin", "application/octet-stream", bytes(range(256)))
    assert text == "" and "unsupported" in how
    text, how = await file_reader.read_file("x.gdoc", "", b"{}")
    assert text == "" and "Drive" in how


def test_any_image_format_becomes_one_the_vision_model_reads():
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (40, 30), "red").save(buf, "BMP")
    raw, mime = file_reader.to_vision_image(buf.getvalue(), "image/bmp")
    assert mime == "image/jpeg" and Image.open(io.BytesIO(raw)).size == (40, 30)
    assert file_reader.to_vision_image(b"png-bytes", "image/png") == (b"png-bytes", "image/png")
    assert file_reader.to_vision_image(b"garbage", "image/heic") is None


def test_shared_google_links_fetch_their_plain_text_export():
    from tools.web.fetch_tool import google_export_url
    assert google_export_url("https://docs.google.com/document/d/abc-1/edit?usp=sharing") == \
        "https://docs.google.com/document/d/abc-1/export?format=txt"
    assert google_export_url("https://docs.google.com/spreadsheets/d/XYZ/edit#gid=0").endswith("/XYZ/export?format=csv")
    assert google_export_url("https://example.com/a") == "https://example.com/a"
